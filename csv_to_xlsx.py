#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
csv_to_xlsx.py — turn the CSV your vision LLM returned into a formatted .xlsx.

Usage:
    python csv_to_xlsx.py llm_output.csv
    python csv_to_xlsx.py llm_output.csv -o venues.xlsx

If the LLM replied with a Markdown table instead of CSV, this script detects it
and converts the table automatically.
"""

import argparse
import csv
import os
import re
import sys

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
except ImportError:
    sys.exit("Missing dependency. Run: pip install -r requirements.txt")

HEADER_FILL = PatternFill(start_color="4C8D4C", end_color="4C8D4C", fill_type="solid")
HEADER_FONT = Font(bold=True, color="FFFFFF")
THIN_BORDER = Border(left=Side(style="thin"), right=Side(style="thin"),
                     top=Side(style="thin"), bottom=Side(style="thin"))


def parse_markdown_table(text):
    """Extract rows from the first Markdown table found in `text`."""
    rows = []
    for raw in text.splitlines():
        line = raw.strip()
        if not line.startswith("|"):
            if rows:
                break
            continue
        cells = [c.strip() for c in line.strip("|").split("|")]
        if all(re.fullmatch(r":?-{2,}:?", c) for c in cells if c):
            continue  # separator row
        rows.append(cells)
    return rows


def parse_csv(text):
    return [row for row in csv.reader(text.splitlines()) if any(c.strip() for c in row)]


def load_rows(path):
    with open(path, encoding="utf-8-sig") as fh:
        text = fh.read()
    if text.lstrip().startswith("|"):
        rows = parse_markdown_table(text)
    else:
        rows = parse_csv(text)
    if not rows:
        sys.exit(f"No data found in {path}")
    return rows


def write_xlsx(rows, out_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "items"

    header, body = rows[0], rows[1:]
    for col, name in enumerate(header, 1):
        cell = ws.cell(row=1, column=col, value=name)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center")
        cell.border = THIN_BORDER

    for r, row in enumerate(body, 2):
        for c, value in enumerate(row, 1):
            cell = ws.cell(row=r, column=c, value=value)
            cell.border = THIN_BORDER

    for col in range(1, len(header) + 1):
        longest = max([len(str(header[col - 1]))] +
                      [len(str(row[col - 1])) for row in body if col <= len(row)] or [10])
        ws.column_dimensions[get_column_letter(col)].width = min(40, max(10, longest + 2))

    ws.freeze_panes = "A2"
    wb.save(out_path)


def main():
    ap = argparse.ArgumentParser(description="Convert LLM CSV/Markdown output to xlsx.")
    ap.add_argument("csv", help="CSV or Markdown file produced by the vision LLM")
    ap.add_argument("-o", "--output", help="output .xlsx path")
    args = ap.parse_args()

    out = args.output or os.path.splitext(args.csv)[0] + ".xlsx"
    rows = load_rows(args.csv)
    write_xlsx(rows, out)
    print(f"Wrote {len(rows) - 1} row(s) -> {out}")


if __name__ == "__main__":
    main()
